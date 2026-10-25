import re, io
from docx import Document
p = r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\thesis\Chapter2 - PSM Assessment System v2.docx"
d = Document(p)
out = io.open("final_check.txt", "w", encoding="utf-8")
txt = "\n".join(x.text for x in d.paragraphs)
w = io.open("final_check.txt", "w", encoding="utf-8")
print("WORDS:", len(txt.split()), file=w)
print(file=w)
cites = sorted(set(re.findall(r"\(([^()]*?(?:19|20)\d\d[^()]*?)\)", txt)))
print("CITATION STRINGS:", len(cites), file=w)
for c in cites: print("   (", c, ")", file=w)
print(file=w)
print("REFERENCES:", file=w)
for x in d.paragraphs:
    if re.match(r"^(i{1,3}|iv|v)\.\s", x.text.strip()):
        print("   ", x.text.strip()[:75], file=w)
print(file=w)
print("HEADINGS:", file=w)
for x in d.paragraphs:
    if x.style.name.startswith("Heading"):
        print("   ", repr(x.text), file=w)
print(file=w)
stale = ["BookIT", "EMS", "MPP", "Event Management", "Room Reservation", "UTMH"]
print("STALE REFERENCES TO OLD SYSTEMS:", file=w)
for s in stale:
    n = txt.count(s) + sum(c.text.count(s) for t in d.tables for r in t.rows for c in r.cells)
    print(f"   {s}: {n}", file=w)
print(file=w)
print("TABLE:", file=w)
t = d.tables[0]
print("   rows:", len(t.rows), "cols:", len(t.columns), file=w)
for r in t.rows: print("   ", [c.text for c in r.cells], file=w)
w.close()
print("ok")
