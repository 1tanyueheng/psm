from docx import Document
import sys

p = r"C:\Users\yuehe\Desktop\study\fyp\thesis\c2\Chapter2.docx"
out = open(r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\.workbuddy-ai\ch2_dump.txt", "w", encoding="utf-8")

d = Document(p)
print("=== PARAGRAPHS:", len(d.paragraphs), "| TABLES:", len(d.tables), "| SECTIONS:", len(d.sections), file=out)
print("=== BODY ORDER ===", file=out)
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

body = d.element.body
idx_p = 0
idx_t = 0
for child in body.iterchildren():
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        para = Paragraph(child, d)
        txt = para.text
        style = para.style.name if para.style is not None else "?"
        n_runs = len(para.runs)
        bold = any(r.bold for r in para.runs)
        print(f"[P{idx_p:03d}] style={style!r} runs={n_runs} bold={bold} :: {txt}", file=out)
        idx_p += 1
    elif tag == 'tbl':
        t = Table(child, d)
        print(f"[TABLE {idx_t}] rows={len(t.rows)} cols={len(t.columns)}", file=out)
        for ri, row in enumerate(t.rows):
            cells = [c.text.replace('\n', ' / ')[:60] for c in row.cells]
            print(f"    r{ri}: {cells}", file=out)
        idx_t += 1
    elif tag == 'sectPr':
        print("[SECTPR]", file=out)
out.close()
print("done")
