import re
from docx import Document
p = r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\thesis\Chapter2 - PSM Assessment System (cited).docx"
d = Document(p)
out = open(r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\.workbuddy-ai\cite_check.txt","w",encoding="utf-8")

body_text = "\n".join(x.text for x in d.paragraphs)
words = len(body_text.split())
in_cites = re.findall(r"\(([^()]*?(?:19|20)\d\d[^()]*?)\)", body_text)
names = set()
for c in in_cites:
    for part in c.split(";"):
        m = re.match(r"\s*([A-Z][A-Za-z .&]+?)[,]?\s*(?:19|20)\d\d", part.strip())
        if m: names.add(m.group(1).strip())
print("In-text citation strings found:", len(in_cites), file=out)
for c in sorted(set(in_cites)):
    print("   (", c, ")", file=out)
print(file=out)
print("Author keys used in text:", sorted(names), file=out)
print(file=out)
print("Reference list:", file=out)
for x in d.paragraphs:
    if re.match(r"^(i{1,3}|iv|v|vi)\.\s", x.text.strip()):
        print("   ", x.text.strip()[:90], file=out)
print(file=out)
print("TOTAL WORDS:", words, file=out)
print(file=out)
print("=== headings ===", file=out)
for x in d.paragraphs:
    if x.style.name.startswith("Heading"):
        print("   ", x.text, file=out)
print(file=out)
print("=== EMS section paragraph count ===", file=out)
paras = [x.text for x in d.paragraphs]
try:
    i = next(k for k,t in enumerate(paras) if t.startswith("Web-Based UTHM Event"))
    j = next(k for k,t in enumerate(paras) if t.startswith("UTHM Student Representative"))
    print("EMS block paragraphs:", j-i, file=out)
    for t in paras[i:j]:
        print("   -", t[:70], file=out)
except StopIteration:
    print("not found", file=out)
out.close()
print("ok")
