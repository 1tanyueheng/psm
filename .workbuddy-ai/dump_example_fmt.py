from docx import Document
from docx.oxml.ns import qn
import re
d = Document(r"C:\Users\yuehe\Downloads\Chapter2.docx")
out = open(r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\.workbuddy-ai\example_fmt.txt","w",encoding="utf-8")

print("=== PARAGRAPH FORMAT (sample of each role) ===", file=out)
for i in [2,5,7,8,9,10,17,18,19,20,21,29,31,33,34,55,59,60,61,62,64,66,67,72,73]:
    p = d.paragraphs[i]
    pf = p.paragraph_format
    r0 = p.runs[0] if p.runs else None
    print(f"[P{i:03d}] style={p.style.name} align={pf.alignment} first={pf.first_line_indent} "
          f"left={pf.left_indent} sb={pf.space_before} sa={pf.space_after} line={pf.line_spacing}", file=out)
    if r0:
        print(f"        run0: font={r0.font.name} size={r0.font.size} bold={r0.bold} italic={r0.italic} | {r0.text[:45]!r}", file=out)
    else:
        print("        (no runs)", file=out)
    print(file=out)

print("=== Normal style ===", file=out)
s = d.styles["Normal"]
print("font:", s.font.name, s.font.size, "| align:", s.paragraph_format.alignment,
      "first:", s.paragraph_format.first_line_indent,
      "line:", s.paragraph_format.line_spacing,
      "sb:", s.paragraph_format.space_before, "sa:", s.paragraph_format.space_after, file=out)
print(file=out)
print("=== Heading 2 / Heading 3 xml ===", file=out)
for n in ["Heading 2","Heading 3"]:
    st = d.styles[n]
    x = re.sub(r'\s+xmlns:[a-z0-9]+="[^"]*"','',st.element.xml)
    print(n, ":", x[:900], file=out)
    print(file=out)

print("=== TABLE xml (first rows) ===", file=out)
t = d.tables[0]
tx = re.sub(r'\s+xmlns:[a-z0-9]+="[^"]*"','',t._tbl.xml)
print(tx[:2600], file=out)
out.close()
print("ok")
