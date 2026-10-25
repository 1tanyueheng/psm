import io
from docx import Document
p = r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\thesis\Chapter4 - Generate Report Use Case (Sequence and Activity Diagrams).docx"
d = Document(p)
w = io.open("rep_check.txt","w",encoding="utf-8")
print("paragraphs:", len(d.paragraphs), "tables:", len(d.tables), "images:", len(d.inline_shapes), file=w)
print("words:", sum(len(x.text.split()) for x in d.paragraphs), file=w)
print(file=w); print("HEADINGS:", file=w)
for x in d.paragraphs:
    if x.style.name.startswith("Heading"):
        print("  ", repr(x.text), file=w)
print(file=w); print("CAPTIONS:", file=w)
for x in d.paragraphs:
    t = x.text.strip()
    if t.startswith(("Table 4.", "Figure 4.")):
        print("  ", t, file=w)
print(file=w); print("SPEC ROWS:", file=w)
for r in d.tables[0].rows:
    print("  ", r.cells[0].text, "->", r.cells[1].text[:44].replace("\n"," | "), file=w)
print(file=w)
bad = [w_ for w_ in ("green","amber","shown in") if w_ in "\n".join(x.text for x in d.paragraphs)]
print("stale colour words:", bad, file=w)
w.close(); print("ok")
