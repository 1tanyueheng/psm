from docx import Document
from docx.oxml.ns import qn
d = Document(r"C:\Users\yuehe\Desktop\study\fyp\thesis\c1\Chapter1.docx")
print("=== Chapter1 heading paras: raw xml snippets ===")
for i, p in enumerate(d.paragraphs[:6]):
    if p.style.name.startswith("Heading"):
        pPr = p._p.find(qn("w:pPr"))
        numpr = pPr.find(qn("w:numPr")) if pPr is not None else None
        print(f"[{i}] {p.style.name} text={p.text!r}")
        print("    numPr:", numpr.xml if numpr is not None else None)
print()
print("=== does Chapter1 have 'CHAPTER 1' typed? ===")
for i, p in enumerate(d.paragraphs[:4]):
    print(i, p.style.name, repr(p.text))
