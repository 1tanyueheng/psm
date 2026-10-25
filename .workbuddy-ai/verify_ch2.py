from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

p = r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\thesis\Chapter2 - Literature Review.docx"
out = open(r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\.workbuddy-ai\ch2_v2.txt","w",encoding="utf-8")
d = Document(p)
print("paragraphs:", len(d.paragraphs), "tables:", len(d.tables), "sections:", len(d.sections), file=out)
print("sections:", [(s.page_width.cm, s.page_height.cm) for s in d.sections][:3], file=out)
print(file=out)
words = 0
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if t:
        words += len(t.split())
        tag = {"Heading 1":"H1","Heading 2":"H2","Heading 3":"H3","Caption":"CAP"}.get(para.style.name, "  ")
        print(f"{tag} [{i:03d}] {t[:110]}", file=out)
print(file=out)
print("BODY WORD COUNT:", words, file=out)
for ti, t in enumerate(d.tables):
    print(file=out)
    print(f"TABLE {ti}: rows={len(t.rows)} cols={len(t.columns)} style={t.style.name if t.style else None}", file=out)
    for row in t.rows:
        print("   ", [c.text[:34] for c in row.cells], file=out)
out.close()
print("ok")
