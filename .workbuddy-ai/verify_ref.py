from docx import Document
p = r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\thesis\Chapter1 - with citations.docx"
d = Document(p)
for i,x in enumerate(d.paragraphs):
    if x.text.strip() == "References":
        print(f"[{i}] style={x.style.name} :: {x.text}")
        print("   next:", d.paragraphs[i+1].text[:60])
print("headings:")
for x in d.paragraphs:
    if x.style.name.startswith("Heading"):
        print("   ", x.style.name, "|", x.text.replace("\n"," ").strip())
