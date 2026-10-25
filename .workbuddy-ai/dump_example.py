from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

src = r"C:\Users\yuehe\Downloads\Chapter2.docx"
out = open(r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\.workbuddy-ai\example_ch2.txt","w",encoding="utf-8")
d = Document(src)
print("paragraphs:", len(d.paragraphs), "tables:", len(d.tables), "sections:", len(d.sections), file=out)
print("sections:", [(round(round(s.page_width)/360000,2), round(round(s.page_height)/360000,2)) for s in d.sections], file=out)
print(file=out)

total_words = 0
body = d.element.body
pi = 0; ti = 0
for child in body.iterchildren():
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        p = Paragraph(child, d)
        t = p.text.strip()
        style = p.style.name if p.style is not None else "?"
        n_runs = len(p.runs)
        bold = [r.text for r in p.runs if r.bold][:1]
        total_words += len(t.split())
        print(f"[P{pi:03d}] style={style!r} runs={n_runs} bold={bold} :: {t}", file=out)
        pi += 1
    elif tag == 'tbl':
        tb = Table(child, d)
        print(f"\n[TABLE {ti}] rows={len(tb.rows)} cols={len(tb.columns)} style={tb.style.name if tb.style else None}", file=out)
        for ri, row in enumerate(tb.rows):
            print(f"    r{ri}: {[c.text.replace(chr(10),' / ') for c in row.cells]}", file=out)
        print(file=out)
        ti += 1
print(file=out)
print("TOTAL WORD COUNT (incl headings):", total_words, file=out)
out.close()
print("ok")
