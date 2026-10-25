"""Build the Chapter 4 UML deliverable: Use Case Specification + Sequence
Diagram + Activity Diagram for "Login and Authentication".

Follows the same visual contract as Chapter 2 (Times New Roman 12pt,
justified, 1.5 spacing, caption style for figures).
"""
import os
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system"
SRC = r"C:\Users\yuehe\Desktop\study\fyp\thesis\c2\Chapter2.docx"
OUT = os.path.join(BASE, r"thesis\Chapter4 - Login Use Case (Sequence and Activity Diagrams).docx")
DIA = os.path.join(BASE, r"thesis\diagrams")

doc = Document(SRC)
body = doc.element.body
sectPr = body.find(qn("w:sectPr"))
import copy
FINAL = copy.deepcopy(list(body.iter(qn("w:sectPr")))[-1])
for ch in list(body):
    if ch is not sectPr:
        body.remove(ch)

TNR = "Times New Roman"


def _run(p, text, bold=False, italic=False, size=12):
    r = p.add_run(text)
    r.font.name = TNR
    r._element.rPr.rFonts.set(qn("w:eastAsia"), TNR)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    return r


def P(text="", bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, italic=False,
      space_after=0):
    p = doc.add_paragraph(style="Normal")
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    p.alignment = align
    if text:
        _run(p, text, bold=bold, italic=italic, size=size)
    return p


_sec = 0


def H(text, level=2):
    global _sec
    p = doc.add_paragraph(style=f"Heading {level}")
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "0")
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.insert(0, numPr)
    if level == 2:
        _sec += 1
        _run(p, f"4.{_sec}\t", bold=True)
    _run(p, text, bold=True)
    pf = p.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(10)
    return p


def caption(text):
    p = P(text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    return p


def figure(path, width_cm, cap):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    caption(cap)
    return p


def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)


def spec_table(rows, label_w=4.2, value_w=10.6):
    t = doc.add_table(rows=0, cols=2)
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    for dup in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(dup)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tblPr.append(tw)
    tw.set(qn("w:type"), "dxa")
    tw.set(qn("w:w"), str(int((label_w + value_w) / 2.54 * 1440)))

    for label, value in rows:
        cells = t.add_row().cells
        trPr = cells[0]._tc.getparent().get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        for i, (txt, bold, fill) in enumerate(((label, True, "E7E6E6"),
                                               (value, False, None))):
            c = cells[i]
            para = c.paragraphs[0]
            para.paragraph_format.line_spacing = 1.15
            para.paragraph_format.space_before = Pt(1)
            para.paragraph_format.space_after = Pt(1)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _run(para, txt, bold=bold, size=10.5)
            if fill:
                shade(c, fill)
        # keep the label and its value on the same page
        for c in cells:
            for pr in c.paragraphs:
                pr.paragraph_format.keep_with_next = True
        cells[0].width = Cm(label_w)
        cells[1].width = Cm(value_w)
    for col, w in enumerate((label_w, value_w)):
        t.columns[col].width = Cm(w)
    return t


# ============================================================
P("CHAPTER 4", bold=True)
P()
P()
P("ANALYSIS AND DESIGN", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
P()
P("Use Case Specification and UML Diagrams", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
P("Login and Authentication", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
P()
P()

# ---------------------------------------------------------- 4.1
H("Introduction")
P("This document presents the detailed design of the first use case of the PSM "
  "Assessment System, namely Login and Authentication. It shows how the use case "
  "identified in the use case diagram is refined into a use case specification, a "
  "sequence diagram and an activity diagram. The use case is described first, then "
  "each diagram is presented and explained step by step.")

# ---------------------------------------------------------- 4.2
H("Use Case Specification")
P("The Login and Authentication use case allows any registered user to sign in. It is "
  "the entry point to every other use case in the system, because all remaining use "
  "cases require an authenticated user whose role determines what the user may access. "
  "Table 4.1 summarises the use case.")

caption("Table 4.1: Use Case Specification for Login and Authentication")
spec_table([
    ("Use Case ID", "UC-01"),
    ("Use Case Name", "Login and Authentication"),
    ("Actor", "Student, Admin, Supervisor, Lecturer, Panel, Coordinator, System Controller"),
    ("Description", "The user signs in with an email address and password. The system "
                    "verifies the credentials, checks that the account is not locked and "
                    "is active, then issues an access token and redirects the user to the "
                    "dashboard for their role."),
    ("Pre-condition", "The user has a registered account in the system, and the account "
                      "has been activated by the administrator."),
    ("Trigger", "The user opens the Login page and submits their email and password."),
    ("Main Flow", "1. The user opens the Login page.\n"
                  "2. The user enters an email address and password.\n"
                  "3. The user clicks Login.\n"
                  "4. The system validates the input format.\n"
                  "5. The system finds the account for that email address.\n"
                  "6. The system verifies the password against the stored hash.\n"
                  "7. The system confirms the account is active.\n"
                  "8. The system creates an access token for the session.\n"
                  "9. The system records the successful login in the audit log.\n"
                  "10. The system returns the token, the user details and the landing "
                  "route for the user's role.\n"
                  "11. The user is redirected to the role dashboard."),
    ("Alternative Flow",
     "A1. Input invalid: at step 4, the system displays an input error message and "
     "returns to step 2.\n"
     "A2. Account not found: at step 5, the system logs the failed attempt and returns "
     "a general authentication failure message, without revealing whether the email is "
     "registered.\n"
     "A3. Account locked: the system returns a message stating that the account is "
     "temporarily locked.\n"
     "A4. Password incorrect: at step 6, the system increases the failed attempt "
     "counter, logs the failure and returns a general authentication failure message. "
     "After five consecutive failures the account is locked for fifteen minutes, and "
     "the lock duration doubles on each further block.\n"
     "A5. Account not active: at step 7, the system returns a message stating that the "
     "account is suspended, inactive or not yet verified."),
    ("Post-condition",
     "On success, the user holds a valid access token and is signed in, and the login "
     "is recorded in the audit log. On failure, no token is issued, the failure is "
     "recorded, and the user remains on the Login page."),
])

P()
P("Several design decisions in Table 4.1 follow good practice for authentication. The "
  "failure messages for an unknown account (A2) and for an incorrect password (A4) are "
  "deliberately identical, so that an attacker cannot use the login form to discover "
  "which email addresses belong to students or staff. The failed attempt counter (A4) "
  "limits password guessing, and the lockout duration increases on each repeat so that "
  "a persistent attacker is slowed progressively rather than merely delayed.")

# ---------------------------------------------------------- 4.3
H("Sequence Diagram")
P("The sequence diagram in Figure 4.1 shows the same flow as an ordered exchange of "
  "messages between the objects involved. Five lifelines are used: the actor, the Login "
  "interface, the authentication controller, the user model with its database, and the "
  "audit logger. The numbering of the messages follows the main flow and the "
  "alternative flows in Table 4.1.")

figure(os.path.join(DIA, "sequence_login.png"), 12.6,
       "Figure 4.1: Sequence Diagram for Login and Authentication")

P("The diagram may be read step by step as follows. Messages 1 and 2 carry the "
  "credentials from the actor, through the Login interface, to the controller. Message "
  "3 validates the input, which corresponds to step 4 of the main flow. Messages 4 and "
  "5 retrieve the account record. Message 6 is the decision that produces alternative "
  "flow A2: when no record is returned, message 7 records the failure and message 8 "
  "returns the general failure message to the interface.")
P("When a record is found, message 9 tests whether the account is locked, which is "
  "alternative flow A3 and is answered by message 10. Messages 11 and 12 verify the "
  "password hash; a mismatch is alternative flow A4, so message 13 increases the "
  "failed attempt counter and applies the lockout rule, message 14 records the "
  "failure, and message 15 returns the same general failure message used for A2. "
  "Message 16 is the final decision, the account status check, which produces "
  "alternative flow A5 through message 17.")
P("Only when all decisions pass does the successful path run. Message 18 creates the "
  "access token, which also resets the failed attempt counter and records the login "
  "time and address, and message 19 returns the token. Message 20 writes the success "
  "record to the audit log. Message 21 returns the response containing the token, the "
  "user details and the landing route for the role, and message 22 redirects the user "
  "to the correct dashboard.")

# ---------------------------------------------------------- 4.4
H("Activity Diagram")
P("The activity diagram in Figure 4.2 presents the same use case as a workflow. It is "
  "useful because it shows clearly where each alternative flow leaves the main path and "
  "where every failure re-enters it, which the sequence diagram expresses less "
  "directly.")

figure(os.path.join(DIA, "activity_login.png"), 14.2,
       "Figure 4.2: Activity Diagram for Login and Authentication")

P("The workflow begins at the start node and moves through the three interface "
  "activities: opening the Login page, entering the email address and password, and "
  "clicking Login. The system then validates the input. The first decision, Valid?, "
  "corresponds to alternative flow A1 and shows the input error message when the input "
  "is rejected.")
P("After the account lookup, the second decision, User found?, corresponds to "
  "alternative flow A2 and records the failed attempt before returning the failure "
  "message. The third decision, Account locked?, corresponds to alternative flow A3. "
  "The fourth decision, Password correct?, corresponds to alternative flow A4; its "
  "failure branch increases the failed attempt counter, applies the lockout rule and "
  "records the failure. The fifth decision, Account active?, corresponds to "
  "alternative flow A5.")
P("When every decision passes, the workflow continues through the success activities: "
  "creating the access token, recording the successful login, returning the response "
  "and redirecting the user to the role dashboard, after which the flow reaches the end "
  "node. The five failure branches are collected on the right hand side of the diagram "
  "and return to the Login form, so that the user may try again.")

# ---------------------------------------------------------- 4.5
H("Chapter Summary")
P("This document has specified the Login and Authentication use case and presented its "
  "sequence diagram and activity diagram. The use case specification defines the actors, "
  "the main flow, the five alternative flows and the pre-conditions and post-conditions. "
  "The sequence diagram shows the same flow as message exchanges between five objects, "
  "and the activity diagram shows it as a workflow with five decision points. The "
  "remaining use cases identified in the use case diagram, namely Manage Assignment, "
  "Manage Milestone, Manage Evaluation, Generate Report, Manage Alert and Notification, "
  "and Manage Title and Registration, follow the same documentation approach.")

body.remove(sectPr)
body.append(FINAL)
doc.save(OUT)
print("saved:", OUT)
