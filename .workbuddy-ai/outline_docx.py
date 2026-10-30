"""Structural read of the thesis deliverables: headings, word counts, tables."""
import sys

from docx import Document

for path in sys.argv[1:]:
    doc = Document(path)
    words = sum(len(p.text.split()) for p in doc.paragraphs)
    print(f"\n=== {path} ===")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} words={words}")
    for p in doc.paragraphs:
        style = p.style.name if p.style is not None else ""
        if style.startswith("Heading") or style.lower().startswith("title"):
            print(f"  [{style}] {p.text.strip()[:110]}")
