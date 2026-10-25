from docx import Document
p = r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\thesis\Chapter1 - with citations.docx"
d = Document(p)
print("paragraphs:", len(d.paragraphs), "tables:", len(d.tables), "sections:", len(d.sections))
import re
full = "\n".join(x.text for x in d.paragraphs)
cites = re.findall(r"\((?:[A-Z][A-Za-z&.,\s]+),\s*(?:19|20)\d\d\)", full)
print("in-text citations:", len(cites))
for c in sorted(set(cites)): print("   ", c)
print("References heading present:", "References" in full)
