from docx import Document
p = r"C:\Users\yuehe\Desktop\study\fyp\thesis\c1\Chapter1.docx"
out = open(r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\.workbuddy-ai\ch1_full.txt", "w", encoding="utf-8")
d = Document(p)
print(f"paragraphs={len(d.paragraphs)} tables={len(d.tables)}", file=out)
for i, para in enumerate(d.paragraphs):
    print(f"[P{i:03d}] {para.style.name} :: {para.text}", file=out)
for ti, t in enumerate(d.tables):
    print(f"\n[TABLE {ti}] rows={len(t.rows)} cols={len(t.columns)}", file=out)
    for ri, row in enumerate(t.rows):
        print(f"   r{ri}: {[c.text.replace(chr(10),' / ')[:70] for c in row.cells]}", file=out)
out.close()
print("ok")
