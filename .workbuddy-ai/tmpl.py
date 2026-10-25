from docx import Document
p = r"C:\Users\yuehe\Desktop\study\fyp\thesis\PSM TESIS.docx"
out = open(r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\.workbuddy-ai\template_dump.txt", "w", encoding="utf-8")
d = Document(p)
print("paras", len(d.paragraphs), "tables", len(d.tables), file=out)
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if t:
        print(f"[P{i:03d}] {para.style.name} :: {t[:160]}", file=out)
out.close()
print("ok")
