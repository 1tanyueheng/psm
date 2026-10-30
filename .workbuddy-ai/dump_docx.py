"""Dump every non-empty paragraph of a .docx with its style, for a close read."""
import sys

from docx import Document

path = sys.argv[1]
doc = Document(path)
print(f"=== {path} ===")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    style = p.style.name if p.style is not None else ""
    print(f"{i:>3} [{style}] {text}")

for ti, table in enumerate(doc.tables):
    print(f"\n--- table {ti} ({len(table.rows)}x{len(table.columns)}) ---")
    for row in table.rows:
        print(" | ".join(c.text.strip().replace("\n", " ")[:60] for c in row.cells))
