import openpyxl, json, sys
p = r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\HIRARC FORM_FSKTMPSM_UTHMOSHEBKPPS_001_Oct 2025.xlsx"
_rf = open(r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\.workbuddy-ai\verify_out.txt", "w", encoding="utf-8")
def print(*a, **k):
    k["file"] = _rf
    return __builtins__.print(*a, **k) if not isinstance(__builtins__, dict) else __builtins__["print"](*a, **k)
wb = openpyxl.load_workbook(p, data_only=False)
ws = wb.active
print("sheets:", wb.sheetnames, "| active:", ws.title)
print("dims:", ws.dimensions, "| merges:", len(ws.merged_cells.ranges))
print("images:", len(ws._images))
print("print_area:", ws.print_area)
print("orientation:", ws.page_setup.orientation, "fitToPage:", ws.sheet_properties.pageSetUpPr.fitToPage,
      "fitW:", ws.page_setup.fitToWidth, "fitH:", ws.page_setup.fitToHeight)
print("gridlines shown:", ws.sheet_view.showGridLines, "| fullCalcOnLoad:", wb.calculation.fullCalcOnLoad)
print("col widths:", {c: round(ws.column_dimensions[c].width, 2) for c in "ABCDEFGHIJK"})
print("row heights:", {r: ws.row_dimensions[r].height for r in range(1, 27)})
print()
print("--- merged ranges ---")
for rng in sorted(ws.merged_cells.ranges, key=lambda x: (x.min_row, x.min_col)):
    v = ws.cell(rng.min_row, rng.min_col).value
    v = (str(v)[:70] + "...") if v and len(str(v)) > 70 else v
    print(f"  {str(rng):12s} {v!r}")
print()
print("--- formulas ---")
for row in ws.iter_rows(min_row=22, max_row=26):
    for c in row:
        if c.column_letter in ("H", "I"):
            print(f"  {c.coordinate}: {c.value}")
print()
print("--- text that overflows its row (est.) ---")
import math
COLW = {c: ws.column_dimensions[c].width for c in "ABCDEFGHIJK"}
def cpl(w, size=8.0):
    return max(4.0, w * (11.0/size) * 0.92)
problems = 0
for row in ws.iter_rows(min_row=1, max_row=26):
    h = ws.row_dimensions[row[0].row].height or 15
    cap = int(h / 10.6)
    for c in row:
        if not isinstance(c.value, str):
            continue
        if c.value.startswith("="):
            continue
        # merged width
        w = COLW[c.column_letter]
        for rng in ws.merged_cells.ranges:
            if c.coordinate == rng.coord.split(":")[0]:
                w = sum(COLW[openpyxl.utils.get_column_letter(i)]
                        for i in range(rng.min_col, rng.max_col + 1))
                break
        size = c.font.size or 11
        n = 0
        for para in c.value.split("\n"):
            n += max(1, math.ceil(len(para) / cpl(w, size))) if para else 1
        if n > cap:
            problems += 1
            print(f"  !! {c.coordinate} needs ~{n} lines, row height {h} fits ~{cap}")
print("overflow problems:", problems)
