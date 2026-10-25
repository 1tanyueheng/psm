from docx import Document
from docx.oxml.ns import qn
for label, path in [("ORIGINAL", r"C:\Users\yuehe\Desktop\study\fyp\thesis\c2\Chapter2.docx"),
                    ("NEW", r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\thesis\Chapter2 - Literature Review.docx")]:
    d = Document(path)
    print("=====", label, "=====")
    print("  sections:", len(d.sections))
    for i, s in enumerate(d.sections):
        hdr = s.header
        ftr = s.footer
        print(f"   sec{i}: {s.page_width.cm:.1f}x{s.page_height.cm:.1f}cm start={s.start_type} "
              f"hdr_linked={hdr.is_linked_to_previous} ftr_linked={ftr.is_linked_to_previous} "
              f"hdr_text={[p.text for p in hdr.paragraphs if p.text.strip()]} "
              f"ftr_text={[p.text for p in ftr.paragraphs if p.text.strip()]}")
    body = d.element.body
    print("  body children tags:", [c.tag.split('}')[-1] for c in body.iterchildren()])
