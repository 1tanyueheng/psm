from docx import Document
d = Document(r"C:\Users\yuehe\Desktop\study\fyp\thesis\c1\Chapter1.docx")
txt = "\n".join(p.text for p in d.paragraphs)
import re
# any existing citations in Ch1?
print("paren yr:", re.findall(r"\([^)]{0,60}(?:19|20)\d\d[^)]{0,40}\)", txt)[:10])
print("has References heading:", any(p.text.strip().lower()=="references" for p in d.paragraphs))
# check the template's References style + an example entry
t = Document(r"C:\Users\yuehe\Desktop\study\fyp\thesis\PSM TESIS.docx")
for i,p in enumerate(t.paragraphs):
    if p.text.strip().lower()=="references":
        for j in range(i+1, min(i+6, len(t.paragraphs))):
            if t.paragraphs[j].text.strip():
                print(f"template ref sample (style={t.paragraphs[j].style.name}):", t.paragraphs[j].text[:180])
        break
