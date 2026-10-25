"""Build Chapter 4 deliverable #3: Manage Milestone use case (UC-03)."""
import copy
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system"
SRC = r"C:\Users\yuehe\Desktop\study\fyp\thesis\c2\Chapter2.docx"
OUT = os.path.join(BASE, r"thesis\Chapter4 - Manage Milestone Use Case (Sequence and Activity Diagrams).docx")
DIA = os.path.join(BASE, r"thesis\diagrams")

doc = Document(SRC)
body = doc.element.body
sectPr = body.find(qn("w:sectPr"))
FINAL = copy.deepcopy(list(body.iter(qn("w:sectPr")))[-1])
for ch in list(body):
    if ch is not sectPr:
        body.remove(ch)

TNR = "Times New Roman"


def _run(p, text, bold=False, size=12):
    r = p.add_run(text)
    r.font.name = TNR
    r._element.rPr.rFonts.set(qn("w:eastAsia"), TNR)
    r.font.size = Pt(size)
    r.bold = bold
    return r


def P(text="", bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, space_after=0):
    p = doc.add_paragraph(style="Normal")
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    p.alignment = align
    if text:
        _run(p, text, bold=bold, size=size)
    return p


_sec = 0


def H(text):
    global _sec
    _sec += 1
    p = doc.add_paragraph(style="Heading 2")
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "0")
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.insert(0, numPr)
    _run(p, f"4.{_sec}\t", bold=True)
    _run(p, text, bold=True)
    pf = p.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(10)
    return p


def caption(text):
    return P(text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)


def figure(path, width_cm, cap):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Cm(width_cm))
    caption(cap)
    return p


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def spec_table(rows, lw=4.2, vw=10.6):
    t = doc.add_table(rows=0, cols=2)
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    for dup in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(dup)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tblPr.append(tw)
    tw.set(qn("w:type"), "dxa")
    tw.set(qn("w:w"), str(int((lw + vw) / 2.54 * 1440)))

    for label, value in rows:
        cells = t.add_row().cells
        for i, (txt, bold, fill) in enumerate(((label, True, "E7E6E6"),
                                               (value, False, None))):
            para = cells[i].paragraphs[0]
            para.paragraph_format.line_spacing = 1.15
            para.paragraph_format.space_before = Pt(1)
            para.paragraph_format.space_after = Pt(1)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _run(para, txt, bold=bold, size=10.5)
            if fill:
                shade(cells[i], fill)
        cells[0].width = Cm(lw)
        cells[1].width = Cm(vw)
    for col, w in enumerate((lw, vw)):
        t.columns[col].width = Cm(w)
    return t


# =====================================================================
P("CHAPTER 4", bold=True)
P()
P()
P("ANALYSIS AND DESIGN", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
P()
P("Use Case Specification and UML Diagrams", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
P("Manage Milestone", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
P()
P()

# ---------------------------------------------------------------- 4.1
H("Introduction")
P("This document presents the detailed design of the third use case of the PSM "
  "Assessment System, namely Manage Milestone. The use case covers the milestone "
  "lifecycle from the moment a project is registered until the milestone is approved "
  "and the next milestone in the chain is opened. It is specified in Table 4.3, then "
  "expressed as a sequence diagram in Figure 4.5 and an activity diagram in Figure 4.6.")

# ---------------------------------------------------------------- 4.2
H("Use Case Specification")
P("A project is created with a fixed chain of milestones generated from the PSM part. "
  "Each milestone moves through a defined set of statuses, and the system refuses any "
  "change that is not a legal transition. The most important rule in the use case is "
  "that approving a milestone is what opens the next one, so the chain cannot run "
  "ahead of the work. Table 4.3 summarises the use case.")

caption("Table 4.3: Use Case Specification for Manage Milestone")
spec_table([
    ("Use Case ID", "UC-03"),
    ("Use Case Name", "Manage Milestone"),
    ("Actor", "Student (submit), Supervisor (review), Coordinator and Admin (review, "
              "deadline, ad-hoc milestone). Panel for the proposal milestone."),
    ("Description", "The system generates a milestone chain for a project and opens the "
                    "milestones in order. A student submits work to the open milestone, "
                    "and a reviewer approves the work or requests a revision. Approving "
                    "a milestone opens the next milestone in the chain."),
    ("Pre-condition", "The project is registered and its milestone chain has been "
                      "generated. The milestone to be submitted is open and within its "
                      "deadline or late window."),
    ("Trigger", "The student uploads work to an open milestone, or a reviewer acts on a "
                "submitted milestone."),
    ("Main Flow",
     "1. The system generates the milestone chain when the project is registered.\n"
     "2. The system opens the first milestone and stamps the opening date.\n"
     "3. The student opens the milestone and uploads one or more files.\n"
     "4. The system validates the file count, size and extensions.\n"
     "5. The system stores the files privately under a randomised name.\n"
     "6. The system records a new revision and marks it as the current version.\n"
     "7. The milestone status becomes submitted and the reviewer is notified.\n"
     "8. The reviewer opens the submission.\n"
     "9. The reviewer chooses Approve or Request Revision.\n"
     "10. The system checks that the requested transition is legal.\n"
     "11. The system updates the status and its timestamps.\n"
     "12. The system records a timeline event and an audit entry.\n"
     "13. On approval, the system opens the next milestone in the chain.\n"
     "14. The system notifies the students of the outcome."),
    ("Alternative Flow",
     "A1. Milestone not open: the student cannot upload because the chain has not "
     "opened this milestone yet.\n"
     "A2. Deadline passed: the upload is refused when the deadline has passed and the "
     "milestone has no late-submission window.\n"
     "A3. Invalid file: the upload is refused when too many files are attached, a file "
     "exceeds the size limit, or a file has an extension the milestone does not accept.\n"
     "A4. Nothing to approve: approval is refused when no submission has been made yet.\n"
     "A5. Proposal milestone: approval through this use case is refused because the "
     "proposal's verdict belongs to the panel, which records the title decision "
     "separately.\n"
     "A6. Illegal transition: the system refuses a status change that is not in the "
     "allowed transition map.\n"
     "A7. Revision requested: the comment is required, the revision counter increases "
     "and the milestone reopens for a new upload; the next milestone is not opened.\n"
     "A8. Withdrawal: a student may withdraw an attached file while the milestone still "
     "accepts work or while it is submitted but not yet reviewed."),
    ("Post-condition",
     "On approval, the milestone is approved with its timestamps recorded, a timeline "
     "event and audit entry exist, the next milestone is open, and the students have "
     "been notified. When no milestones remain outstanding the project is marked "
     "completed. On a revision request, the milestone is reopened for a new upload and "
     "the chain does not advance. On failure, the status is unchanged."),
])

P()
P("Two rules in Table 4.3 shape the rest of the system. First, the proposal milestone "
  "is deliberately excluded from the generic approval path (A5). Its verdict is the "
  "panel's and it settles the project title, so allowing a supervisor to approve it "
  "here would close the milestone without ever recording a title decision. Second, a "
  "student may still upload after approval, because the faculty rule is that a student "
  "may always improve their work; the new upload supersedes the approved file while "
  "leaving the original on the record, so the timeline still shows what was approved "
  "and when.")

# ---------------------------------------------------------------- 4.3
H("Sequence Diagram")
P("The sequence diagram in Figure 4.5 shows the approval path, which is the part of "
  "the use case that changes the state of more than one milestone. Six lifelines take "
  "part: the reviewer, the milestone interface, the milestone controller, the milestone "
  "model with its database, the milestone service, and the notification component.")

figure(os.path.join(DIA, "sequence_milestone.png"), 12.8,
       "Figure 4.5: Sequence Diagram for Manage Milestone")

P("Messages 1 and 2 carry the reviewer's decision to the controller. Message 3 checks "
  "that the reviewer is the student's own supervisor, a coordinator or an administrator, "
  "and message 4 checks that there is a submission to approve. If the milestone has no "
  "submitted work the system refuses the request, because approving a document that was "
  "never filed is not evidence of anything.")
P("Message 5 is the guard that excludes the proposal milestone. Once it passes, message "
  "7 hands the change to the milestone service, which owns the state machine. Message 8 "
  "checks the requested transition against the allowed transition map, so an illegal "
  "change such as pending to approved is refused rather than written. Messages 10 and 11 "
  "update the status with its timestamps inside a transaction, and messages 12 and 13 "
  "record the timeline event and the audit entry.")
P("Messages 14 to 19 are the chain rule. Message 14 asks the service to open the next "
  "milestone after this one, selecting the next pending, overdue or rejected milestone "
  "in sequence order. Message 17 opens it and stamps its opening date, and when the "
  "deadline is relative the service also moves the deadline so the milestone does not "
  "open already overdue. If no milestone remains, message 19 marks the project "
  "completed. Messages 21 to 23 notify the students and return the updated milestone.")

# ---------------------------------------------------------------- 4.4
H("Activity Diagram")
P("The activity diagram in Figure 4.6 shows the whole lifecycle, from the creation of "
  "the chain to the end of the milestone. It is the clearer of the two diagrams for "
  "explaining the status model, because it shows both outcomes of the review side by "
  "side.")

figure(os.path.join(DIA, "activity_milestone.png"), 14.2,
       "Figure 4.6: Activity Diagram for Manage Milestone")

P("The workflow begins when the coordinator registers a project, which generates the "
  "milestone chain and opens the first milestone. The student then uploads work, and "
  "the first decision asks whether the milestone accepts a submission. This corresponds "
  "to alternative flows A1 and A2: a milestone that the chain has not opened yet, or one "
  "whose deadline has passed without a late window, refuses the upload.")
P("After a successful upload the milestone becomes submitted and the reviewer opens it. "
  "The second decision is the reviewer's choice, and the two branches are the core of "
  "the use case. On the Approve branch, which runs down the right of the diagram, the transition is checked, the "
  "proposal milestone is excluded, the status and timestamps are written, the timeline "
  "event and audit entry are recorded, and the next milestone is opened before the "
  "students are notified. On the Request Revision branch, which continues down the centre, the comment is "
  "required, the revision counter increases, the status becomes rejected and the students "
  "are notified. That branch returns to the upload step, which is how the same milestone "
  "can be submitted again without the chain advancing.")
P("Both branches then converge and the workflow reaches the end node. The rule that only "
  "approval opens the next milestone is visible in the diagram as the only route from the "
  "review decision to the chain-activation activity.")

# ---------------------------------------------------------------- 4.5
H("Chapter Summary")
P("This document has specified the Manage Milestone use case and presented its sequence "
  "diagram and activity diagram. The specification defines the actors, a fourteen step "
  "main flow, eight alternative flows and the post-conditions for approval and for a "
  "revision request. The sequence diagram shows the approval path with the chain rule, "
  "and the activity diagram shows the complete lifecycle with both review outcomes and "
  "the re-upload loop.")

body.remove(sectPr)
body.append(FINAL)
doc.save(OUT)
print("saved:", OUT)
