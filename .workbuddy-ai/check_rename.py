from docx import Document
p = r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\thesis\Chapter2 - PSM Assessment System.docx"
out = open(r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\.workbuddy-ai\rename_check.txt","w",encoding="utf-8")
d = Document(p)
print("=== occurrences of 'Assessment System' ===", file=out)
for i, para in enumerate(d.paragraphs):
    if "Assessment System" in para.text:
        print(f"[P{i:03d}] {para.style.name}: {para.text}", file=out)
print(file=out)
print("=== occurrences of 'Management System' (should ONLY be the MPP system) ===", file=out)
for i, para in enumerate(d.paragraphs):
    if "Management System" in para.text:
        print(f"[P{i:03d}] {para.text[:150]}", file=out)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            if "Assessment System" in c.text or "Management System" in c.text:
                print(f"[TABLE] {c.text!r}", file=out)
print(file=out)
print("=== headings ===", file=out)
for i, para in enumerate(d.paragraphs):
    if para.style.name.startswith("Heading"):
        print(f"[H{i:03d}] {para.text}", file=out)
out.close()
print("ok")
