from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
p = r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\thesis\Chapter2 - Literature Review (v3).docx"
out = open(r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\.workbuddy-ai\v3_verify.txt","w",encoding="utf-8")
d = Document(p)
words = 0
print(f"paragraphs={len(d.paragraphs)} tables={len(d.tables)} sections={len(d.sections)}", file=out)
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if t:
        words += len(t.split())
        st = para.style.name
        bold = all(r.bold for r in para.runs if r.text.strip()) if para.runs else False
        print(f"[{i:03d}] {st:12s} bold={bold!s:5s} :: {t[:95]}", file=out)
print(file=out)
print("TOTAL WORDS:", words, file=out)
for t in d.tables:
    print(f"\nTABLE rows={len(t.rows)} cols={len(t.columns)} style={t.style.name}", file=out)
    for row in t.rows:
        print("   ", [c.text for c in row.cells], file=out)
out.close()
print("ok")
