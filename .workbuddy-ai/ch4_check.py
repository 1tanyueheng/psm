import io
from docx import Document
p = r"C:\Users\yuehe\Desktop\study\fyp\AIProject\psm-system\thesis\Chapter4 - Login Use Case (Sequence and Activity Diagrams).docx"
d = Document(p)
w = io.open("ch4_check.txt", "w", encoding="utf-8")
print("paragraphs:", len(d.paragraphs), "tables:", len(d.tables), file=w)
print("images:", len(d.inline_shapes), file=w)
words = sum(len(x.text.split()) for x in d.paragraphs)
print("words:", words, file=w)
print(file=w)
print("HEADINGS:", file=w)
for x in d.paragraphs:
    if x.style.name.startswith("Heading"):
        print("  ", repr(x.text), file=w)
print(file=w)
print("CAPTIONS / FIGURES:", file=w)
for x in d.paragraphs:
    t = x.text.strip()
    if t.startswith(("Table 4.", "Figure 4.")):
        print("  ", t, file=w)
print(file=w)
print("SPEC TABLE ROWS:", file=w)
for r in d.tables[0].rows:
    print("  ", r.cells[0].text, "->", r.cells[1].text[:60].replace("\n", " | "), file=w)
w.close()
print("ok")
